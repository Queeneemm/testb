import json
import logging
import os
import re
from datetime import date, datetime, time
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from telegram import Document, Update
from telegram.ext import Application, CommandHandler, ContextTypes, MessageHandler, filters
from zoneinfo import ZoneInfo

load_dotenv()

logging.basicConfig(format="%(asctime)s %(levelname)s %(name)s: %(message)s", level=logging.INFO)
LOGGER = logging.getLogger(__name__)

DATE_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})$")
HEADERS = ("ФИО", "Дата")
DATA_FILE = Path(os.getenv("DATA_FILE", "data/birthdays.xlsx"))
STATE_FILE = Path(os.getenv("STATE_FILE", "data/sent_notifications.json"))
TIMEZONE = ZoneInfo(os.getenv("TIMEZONE", "Europe/Moscow"))
REMINDER_TIME = os.getenv("REMINDER_TIME", "09:00")


def parse_admin_ids() -> set[int]:
    raw = os.getenv("ADMIN_IDS", "")
    return {int(item.strip()) for item in raw.split(",") if item.strip()}


ADMIN_IDS = parse_admin_ids()


def ensure_data_file() -> None:
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    if not DATA_FILE.exists():
        save_birthdays([])


def parse_birthday(value: object) -> str:
    if isinstance(value, datetime):
        value = value.strftime("%d.%m")
    elif isinstance(value, date):
        value = value.strftime("%d.%m")
    else:
        value = str(value or "").strip()

    match = DATE_RE.match(value)
    if not match:
        raise ValueError(f"дата '{value}' должна быть в формате ДД.ММ")

    day = int(match.group(1))
    month = int(match.group(2))
    try:
        date(2000, month, day)
    except ValueError as exc:
        raise ValueError(f"дата '{value}' не существует") from exc
    return f"{day:02d}.{month:02d}"


def read_birthdays(path: Path = DATA_FILE) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    if sheet.max_column > 2:
        for row in sheet.iter_rows(min_col=3, values_only=True):
            if any(value is not None for value in row):
                raise ValueError("таблица должна содержать только два столбца: ФИО и Дата")

    headers = [str(sheet.cell(row=1, column=col).value or "").strip() for col in (1, 2)]
    if headers != list(HEADERS):
        raise ValueError("первая строка должна содержать заголовки: ФИО, Дата")

    rows: list[dict[str, str]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        name_raw, birthday_raw = row
        if name_raw is None and birthday_raw is None:
            continue
        name = str(name_raw or "").strip()
        if not name:
            raise ValueError(f"строка {row_number}: ФИО не может быть пустым")
        rows.append({"name": name, "birthday": parse_birthday(birthday_raw)})
    return rows


def save_birthdays(rows: Iterable[dict[str, str]], path: Path = DATA_FILE) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Birthdays"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row["name"], row["birthday"]])
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["B"].width = 12
    workbook.save(path)


def load_sent_state() -> set[str]:
    if not STATE_FILE.exists():
        return set()
    return set(json.loads(STATE_FILE.read_text(encoding="utf-8")))


def save_sent_state(state: set[str]) -> None:
    STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
    STATE_FILE.write_text(json.dumps(sorted(state), ensure_ascii=False, indent=2), encoding="utf-8")


def is_allowed(update: Update) -> bool:
    user = update.effective_user
    return bool(user and user.id in ADMIN_IDS)


def restricted(handler):
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not is_allowed(update):
            user_id = update.effective_user.id if update.effective_user else "unknown"
            LOGGER.warning("Denied access for Telegram ID %s", user_id)
            if update.effective_message:
                await update.effective_message.reply_text(f"Нет доступа. Ваш Telegram ID: {user_id}")
            return
        await handler(update, context)
    return wrapper


HELP_TEXT = """Команды:
/start или /help — справка
/upload — как загрузить Excel
/export — выгрузить текущую таблицу
/list — показать дни рождения текстом

Чтобы обновить таблицу, просто отправьте .xlsx файл с двумя столбцами: ФИО и Дата."""


@restricted
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await update.effective_message.reply_text(HELP_TEXT)


@restricted
async def upload_help(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await update.effective_message.reply_text("Отправьте .xlsx файл с заголовками 'ФИО' и 'Дата'. Дата: ДД.ММ, например 25.10.")


@restricted
async def export_table(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    ensure_data_file()
    await update.effective_message.reply_document(document=DATA_FILE.open("rb"), filename="birthdays.xlsx")


@restricted
async def list_birthdays(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    ensure_data_file()
    rows = read_birthdays()
    if not rows:
        await update.effective_message.reply_text("Таблица пока пустая.")
        return
    text = "\n".join(f"{item['name']} — {item['birthday']}" for item in rows)
    await update.effective_message.reply_text(text[:4000])


@restricted
async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    document: Document = update.effective_message.document
    if not document.file_name.lower().endswith(".xlsx"):
        await update.effective_message.reply_text("Нужен файл в формате .xlsx")
        return

    telegram_file = await document.get_file()
    with NamedTemporaryFile(suffix=".xlsx") as tmp:
        await telegram_file.download_to_drive(tmp.name)
        try:
            rows = read_birthdays(Path(tmp.name))
        except Exception as exc:
            await update.effective_message.reply_text(f"Не удалось загрузить таблицу: {exc}")
            return
    save_birthdays(rows)
    await update.effective_message.reply_text(f"Таблица обновлена. Записей: {len(rows)}")


def date_for_year(day: int, month: int, year: int) -> date:
    try:
        return date(year, month, day)
    except ValueError:
        if day == 29 and month == 2:
            return date(year, 3, 1)
        raise


def next_date_for(day_month: str, today: date) -> date:
    day, month = map(int, day_month.split("."))
    candidate = date_for_year(day, month, today.year)
    if candidate < today:
        candidate = date_for_year(day, month, today.year + 1)
    return candidate


async def send_reminders(context: ContextTypes.DEFAULT_TYPE) -> None:
    ensure_data_file()
    today = datetime.now(TIMEZONE).date()
    sent = load_sent_state()
    changed = False

    for item in read_birthdays():
        birthday_date = next_date_for(item["birthday"], today)
        days_left = (birthday_date - today).days
        if days_left not in (0, 10):
            continue

        key = f"{birthday_date.year}:{item['name']}:{item['birthday']}:{days_left}"
        if key in sent:
            continue

        if days_left == 10:
            message = f"У {item['name']} день рождения через 10 дней ({item['birthday']})."
        else:
            message = f"У {item['name']} сегодня день рождения! ({item['birthday']})"

        for chat_id in ADMIN_IDS:
            await context.bot.send_message(chat_id=chat_id, text=message)
        sent.add(key)
        changed = True

    if changed:
        save_sent_state(sent)


def parse_reminder_time(value: str) -> time:
    hours, minutes = map(int, value.split(":", maxsplit=1))
    return time(hour=hours, minute=minutes, tzinfo=TIMEZONE)


def main() -> None:
    if not os.getenv("BOT_TOKEN"):
        raise RuntimeError("BOT_TOKEN is required")
    if not ADMIN_IDS:
        raise RuntimeError("ADMIN_IDS is required")
    ensure_data_file()

    application = Application.builder().token(os.environ["BOT_TOKEN"]).build()
    application.add_handler(CommandHandler(["start", "help"], start))
    application.add_handler(CommandHandler("upload", upload_help))
    application.add_handler(CommandHandler("export", export_table))
    application.add_handler(CommandHandler("list", list_birthdays))
    application.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    application.job_queue.run_daily(send_reminders, time=parse_reminder_time(REMINDER_TIME), name="birthday-reminders")
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
